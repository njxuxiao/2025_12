import pandas as pd
import os
import sys
from openpyxl import load_workbook
import tempfile
import time
from zipfile import BadZipFile
import requests
import urllib.request
import io

try:
    import keyboard
except ImportError:
    print("Error: 'keyboard' library is missing.")
    print("Please run in your terminal or command line: pip install keyboard")
    sys.exit()

try:
    import pykakasi
except ImportError:
    print("Error: 'pykakasi' library is missing.")
    print("Please run in your terminal or command line: pip install pykakasi")
    sys.exit()

try:
    from gtts import gTTS
    import pygame
    gtts_available = True
except ImportError:
    print("Warning: gTTS or pygame library not installed. Online TTS function will be unavailable.")
    print("  - pip install gTTS")
    print("  - pip install pygame")
    gtts_available = False

try:
    import pyttsx3
    pyttsx3_available = True
except ImportError:
    print("Warning: pyttsx3 library not installed. Offline backup TTS function will be unavailable.")
    print("  - pip install pyttsx3")
    pyttsx3_available = False

# Global flags to optimize and silence repetitive TTS logging
_gtts_connection_info_printed = False
_gtts_successful_tld = None

def get_pyttsx3_japanese_voice_id():
    if not pyttsx3_available:
        return None
    try:
        temp_engine = pyttsx3.init()
        voices = temp_engine.getProperty('voices')
        temp_engine.stop()
        del temp_engine

        japanese_voice_id = None
        for voice in voices:
            lang_str = getattr(voice, 'lang', '').lower()
            name_str = getattr(voice, 'name', '').lower()
            if 'ja' in lang_str or 'japanese' in name_str:
                japanese_voice_id = voice.id
                break
        
        if japanese_voice_id:
            print("\nSuccessfully found [Local] Japanese voice ID. Offline backup is available.")
            return japanese_voice_id
        else:
            print("\nWarning: Japanese [Local] voice package not found in your system. Offline backup may not work.")
            return None
    except Exception as e:
        print(f"\nFailed to scan for local voices: {e}")
        return None


def speak_with_gtts(text):
    global _gtts_connection_info_printed, _gtts_successful_tld
    if not gtts_available or not text:
        return False
    
    success = False
    tld_list = [_gtts_successful_tld] if _gtts_successful_tld else ['co.jp', 'com', 'ca', 'com.au']
    proxy_address = 'http://127.0.0.1:7890'
    
    original_http_proxy = os.environ.get('HTTP_PROXY')
    original_https_proxy = os.environ.get('HTTPS_PROXY')

    try:
        os.environ['HTTP_PROXY'] = proxy_address
        os.environ['HTTPS_PROXY'] = proxy_address
        if not _gtts_connection_info_printed:
            print(f"  -> Temporarily set system proxy for gTTS: {proxy_address}")

        audio_buffer = io.BytesIO()

        for tld in tld_list:
            try:
                if not _gtts_connection_info_printed:
                    print(f"  -> Trying Google TTS server via tld='{tld}'...")
                
                tts = gTTS(text=text, lang='ja', tld=tld)
                tts.write_to_fp(audio_buffer)
                success = True
                
                if not _gtts_connection_info_printed:
                    print(f"     Success! Connected via '{tld}'. Subsequent messages will be silenced.")
                    _gtts_connection_info_printed = True
                    _gtts_successful_tld = tld
                break 
            except Exception as e:
                if not _gtts_connection_info_printed:
                    print(f"     Connection to '{tld}' failed: {e}")
                continue

        if success:
            audio_buffer.seek(0)
            pygame.mixer.stop()
            sound = pygame.mixer.Sound(audio_buffer)
            sound.play()
        return success
    except Exception as e:
        print(f"\n!! An unknown error occurred during speech synthesis: {e}")
        return False
    finally:
        if original_http_proxy is None:
            if 'HTTP_PROXY' in os.environ: del os.environ['HTTP_PROXY']
        else:
            os.environ['HTTP_PROXY'] = original_http_proxy
        if original_https_proxy is None:
            if 'HTTPS_PROXY' in os.environ: del os.environ['HTTPS_PROXY']
        else:
            os.environ['HTTPS_PROXY'] = original_https_proxy


def speak_with_pyttsx3(voice_id, text):
    if voice_id and text:
        try:
            # Forcing stop for any lingering pygame sounds when switching to offline
            if gtts_available:
                pygame.mixer.stop()
            engine = pyttsx3.init()
            engine.setProperty('voice', voice_id)
            engine.say(text)
            engine.runAndWait()
            engine.stop()
        except Exception as e:
            print(f"\n!! Error using local TTS: {e}")


def get_display_length(s):
    length = 0
    for char in str(s):
        if '\u4e00' <= char <= '\u9fff':
            length += 2
        else:
            length += 1
    return length

def display_term(word, grammar):
    print("\n" + "╔" + "═"*50 + "╗")
    print("║" + " "*50 + "║")
    
    if word:
        line = f"  [Word]: {word}"
        padding = 50 - get_display_length(line)
        if padding < 0: padding = 0
        print("║" + line + " "*padding + "║")
        
    if grammar:
        line = f"  [Grammar]: {grammar}"
        padding = 50 - get_display_length(line)
        if padding < 0: padding = 0
        print("║" + line + " "*padding + "║")

    print("║" + " "*50 + "║")
    print("╚" + "═"*50 + "╝")

def display_details(meaning, remarks):
    if not meaning and not remarks:
        return

    print("╭" + "┈"*50 + "╮")
    
    if meaning:
        line = f"  [Meaning]: {meaning}"
        padding = 50 - get_display_length(line)
        if padding < 0: padding = 0
        print("┆" + line + " "*padding + "┆")

    if remarks:
        line = f"  [Remarks]: {remarks}"
        padding = 50 - get_display_length(line)
        if padding < 0: padding = 0
        print("┆" + line + " "*padding + "┆")

    print("╰" + "┈"*50 + "╯")


def study_helper(file_path, sheet_to_study=None, tts_mode='auto', tts_timing_mode='immediate'):
    if gtts_available:
        pygame.init()
        pygame.mixer.init()
    
    try:
        japanese_voice_id = get_pyttsx3_japanese_voice_id()
        auto_mode_current_engine = 'gTTS' if gtts_available else 'pyttsx3'

        # tts_timing_mode is now passed as an argument. 
        # Default is 'immediate', can be set to 'after_answer' in main.

        kks = pykakasi.kakasi()

        all_sheets_data = {}
        chosen_sheet = None
        try:
            if not os.path.exists(file_path):
                print(f"Error: File not found '{file_path}'.")
                return

            if not file_path.endswith('.xlsx'):
                print(f"Error: This feature requires an .xlsx format Excel file.")
                print(f"Please open '{os.path.basename(file_path)}' with Excel and save it as .xlsx format.")
                return
            
            xls = pd.ExcelFile(file_path, engine='openpyxl')
            sheet_names = xls.sheet_names

            if not sheet_names:
                print("Error: No worksheets found in the Excel file.")
                return

            if sheet_to_study is not None:
                if isinstance(sheet_to_study, int):
                    if 0 <= sheet_to_study < len(sheet_names):
                        chosen_sheet = sheet_names[sheet_to_study]
                    else:
                        print(f"Error: Specified sheet index {sheet_to_study} is invalid. Valid range is 0 to {len(sheet_names)-1}.")
                        return
                elif isinstance(sheet_to_study, str):
                    if sheet_to_study in sheet_names:
                        chosen_sheet = sheet_to_study
                    else:
                        print(f"Error: Cannot find worksheet named '{sheet_to_study}'.")
                        print(f"Available worksheets are: {sheet_names}")
                        return
                else:
                    print("Error: Invalid type for sheet_to_study parameter. It should be an integer (index) or a string (name).")
                    return
                print(f"Selected worksheet as specified: '{chosen_sheet}'")
            else:
                if len(sheet_names) == 1:
                    chosen_sheet = sheet_names[0]
                    print(f"Automatically selected the only worksheet: '{chosen_sheet}'")
                else:
                    print("Multiple worksheets (Sheets) found:")
                    for i, name in enumerate(sheet_names):
                        print(f"  {i+1}: {name}")
                    while True:
                        try:
                            choice = int(input(f"Please enter the number of the worksheet you want to study (1-{len(sheet_names)}): "))
                            if 1 <= choice <= len(sheet_names):
                                chosen_sheet = sheet_names[choice-1]
                                break
                            else:
                                print("Invalid number, please try again.")
                        except ValueError:
                            print("Please enter a number.")
            
            all_sheets_data = pd.read_excel(file_path, sheet_name=None, engine='openpyxl')
            df = all_sheets_data[chosen_sheet]

            df.columns = df.columns.str.strip()

        except BadZipFile:
            print(f"\nError: The file '{os.path.basename(file_path)}' seems to be corrupted or is not a valid .xlsx file.")
            print("This usually happens if the file was saved incorrectly or is an old .xls file renamed to .xlsx.")
            print("\n[SOLUTION]: Please open the file in Microsoft Excel, go to 'File' -> 'Info' -> and click the 'Convert' button. Then save the file.")
            print("If you don't see a 'Convert' button, using 'File' -> 'Save As' and choosing 'Excel Workbook (*.xlsx)' will also fix it.")
            return
        except Exception as e:
            print(f"Error reading or selecting worksheet: {e}")
            return

        if 'Fre' not in df.columns:
            print("No 'Fre' column detected, creating it automatically.")
            df['Fre'] = 0
        else:
            df['Fre'] = pd.to_numeric(df['Fre'], errors='coerce').fillna(0).astype(int)

        if '单词' not in df.columns and '文法' not in df.columns:
            print(f"Error: Worksheet '{chosen_sheet}' must contain at least a '单词' or '文法' column.")
            return
            
        has_reading_col = '读音' in df.columns
        has_meaning_col = '含义' in df.columns
        has_remarks_col = '备注' in df.columns
        if not has_reading_col:
            print("Info: No '读音' (Reading) column in your Excel. Readings will be auto-generated.")
        if not has_meaning_col:
            print("Info: No '含义' (Meaning) column in your Excel, definitions will not be shown.")
        if not has_remarks_col:
            print("Info: No '备注' (Remarks) column in your Excel, remarks will not be shown.")

        df.sort_values(by='Fre', ascending=False, inplace=True)
        print("\nSorted by 'Fre' (Frequency). The most forgotten items will appear first.")

        print("\n--- Japanese Study Helper Started ---")
        if tts_mode == 'online':
            print("[TTS Mode]: Online Only")
        elif tts_mode == 'offline':
            print("[TTS Mode]: Offline Only")
        else:
            print("[TTS Mode]: Auto (Online first, fallback to Offline)")
        
        print(f"[TTS Timing]: {tts_timing_mode} (Press 't' to toggle)")
            
        print("[IMPORTANT] Please make sure your input method is in English mode to ensure key presses are registered correctly.")
        
        is_changed = False
        last_answered_correctly_index = None
        records = df.to_dict('records')
        original_indices = df.index.tolist()

        i = 0
        enter_press_count = 0
        while i < len(records):
            current_record = records[i]
            original_index = original_indices[i]

            word = str(current_record.get('单词', '')) if pd.notna(current_record.get('单词')) else ""
            grammar = str(current_record.get('文法', '')) if pd.notna(current_record.get('文法')) else ""
            
            reading = ""
            if has_reading_col and pd.notna(current_record.get('读音')):
                reading = str(current_record.get('读音'))
            elif word:
                try:
                    result = kks.convert(word)
                    reading = "".join([item['hira'] for item in result])
                except Exception as e:
                    print(f"!! Could not convert '{word}' to hiragana: {e}")

            meaning = str(current_record.get('含义', '')) if has_meaning_col and pd.notna(current_record.get('含义')) else ""
            remarks = str(current_record.get('备注', '')) if has_remarks_col and pd.notna(current_record.get('备注')) else ""

            if not word and not grammar:
                i += 1
                continue
            
            display_term(word, grammar)

            text_to_speak = word if word else grammar
            
            def speak():
                if tts_mode == 'online':
                    if not speak_with_gtts(text_to_speak):
                        print("\n!! Online TTS failed.")
                    return

                if tts_mode == 'offline':
                    if japanese_voice_id:
                        speak_with_pyttsx3(japanese_voice_id, text_to_speak)
                    else:
                        print("\n!! Offline TTS is not available.")
                    return
                
                nonlocal auto_mode_current_engine
                if auto_mode_current_engine == 'gTTS':
                    if not speak_with_gtts(text_to_speak):
                        print("\n!! Online TTS failed, automatically switching to [Offline TTS] mode.")
                        auto_mode_current_engine = 'pyttsx3'
                        speak_with_pyttsx3(japanese_voice_id, text_to_speak)
                elif auto_mode_current_engine == 'pyttsx3':
                    speak_with_pyttsx3(japanese_voice_id, text_to_speak)

            # --- Logic for Immediate TTS ---
            if tts_timing_mode == 'immediate':
                speak()
            # -------------------------------

            key = None
            while True:
                # Update Prompt to show current TTS Timing Mode
                tts_status = "Immediate" if tts_timing_mode == 'immediate' else "After Ans"
                prompt = f"Action: [→:Know | 0:Forget | t:TTS({tts_status}) | q:Quit"

                if last_answered_correctly_index is not None:
                    prompt += " | x:Undo]"
                else:
                    prompt += "]"

                # Using carriage return \r to overwrite the prompt line if loop continues (cleaner UI)
                print(f"\r{prompt} Progress: {i+1}/{len(records)} ", end="", flush=True)
                
                event = keyboard.read_event(suppress=True)
                while event.event_type != keyboard.KEY_DOWN:
                    event = keyboard.read_event(suppress=True)
                
                key = event.name.lower()

                # --- Toggle TTS Mode ---
                if key == 't':
                    if tts_timing_mode == 'immediate':
                        tts_timing_mode = 'after_answer'
                    else:
                        tts_timing_mode = 'immediate'
                    # Clear line to refresh prompt
                    print(" " * 100, end="\r") 
                    continue
                # -----------------------

                if key == 'enter':
                    enter_press_count += 1
                    if enter_press_count > 3:
                        os.system('cls' if os.name == 'nt' else 'clear')
                        print("\n" * 2)
                        display_term(word, grammar)
                    else:
                        print("\n" * 2)
                        display_term(word, grammar)
                    continue
                
                enter_press_count = 0
                break
            
            print() # Newline after the \r prompt
            
            if key == 'q':
                print("Saving progress and exiting...")
                break 
            
            if key == 'x':
                if last_answered_correctly_index is not None:
                    for idx, record in enumerate(records):
                        if original_indices[idx] == last_answered_correctly_index:
                            record['Fre'] += 1
                            break
                    is_changed = True
                    print(f"Corrected the previous item!")
                    last_answered_correctly_index = None
                else:
                    print("There is no previous item to correct.")
                continue

            last_answered_correctly_index = None
            display_remarks = remarks if remarks else reading

            if key == '0':
                current_record['Fre'] += 1
                is_changed = True
                print(f"Recorded! Forgotten count: {current_record['Fre']}")
                display_details(meaning, display_remarks)
                
                # --- Logic for After Answer TTS ---
                if tts_timing_mode == 'after_answer':
                    speak()
                # ----------------------------------
                
                time.sleep(2)

            elif key == 'right':
                display_details(meaning, display_remarks)
                print(f"Great! Forgotten count: {current_record['Fre']}")
                
                # --- Logic for After Answer TTS ---
                if tts_timing_mode == 'after_answer':
                    speak()
                # ----------------------------------

                last_answered_correctly_index = original_index

            i += 1

        print("\nAll words/grammar have been studied!")

        if not is_changed:
            print("\nNo changes were made, no need to save.")
            return
            
        try:
            print("Updating records back to DataFrame...")
            new_df = pd.DataFrame(records) 
            if 'Fre' in new_df.columns:
                print("Re-sorting by 'Fre' before saving...")
                new_df.sort_values(by='Fre', ascending=False, inplace=True)
            all_sheets_data[chosen_sheet] = new_df

            print("Saving file and preserving column widths...")
            book = load_workbook(file_path)
            col_widths = {}
            for sheet_name in book.sheetnames:
                col_widths[sheet_name] = {
                    letter: dim.width for letter, dim in book[sheet_name].column_dimensions.items()
                }
            
            with pd.ExcelWriter(file_path, engine='openpyxl') as writer:
                for sheet_name, sheet_data in all_sheets_data.items():
                    sheet_data.to_excel(writer, sheet_name=sheet_name, index=False)
                    
                    if sheet_name in col_widths:
                        ws = writer.sheets[sheet_name]
                        for col_letter, width in col_widths[sheet_name].items():
                            if width:
                               ws.column_dimensions[col_letter].width = width

            print("\nStudy session finished! Your progress has been saved successfully, and column widths are preserved.")
        except PermissionError:
            print(f"\nError saving file: Permission denied. Please close the Excel file '{file_path}' and try again.")
        except Exception as e:
            print(f"\nAn unknown error occurred while saving the file: {e}")
    finally:
        if gtts_available:
            pygame.quit()


if __name__ == '__main__':
    
    excel_list = [f for f in os.listdir('.') if f.endswith('.xlsx')]
    target_file = 'ep1_1.xlsx'
    if target_file in excel_list:
        excel_list.remove(target_file)
        has_all = True
    else:
        has_all = False

    excel_list.sort(key=lambda x: int(x.split('_')[0]))

    if has_all:
        excel_list.append(target_file)

###############################################
    if len(excel_list) >= 2:
        excel_file_path = excel_list[-2]
    elif len(excel_list) > 0:
        excel_file_path = excel_list[0]
    else:
        print("No .xlsx files found.")
        sys.exit()

    study_sheet = 1
###############################################
    preferred_tts_engine = 'auto' 
    
    # Options: 'immediate' (speak when word appears) OR 'after_answer' (speak after you press know/forget)
    initial_timing_mode = 'after_answer' 

    study_helper(
        excel_file_path, 
        sheet_to_study=study_sheet, 
        tts_mode=preferred_tts_engine, 
        tts_timing_mode=initial_timing_mode
    )