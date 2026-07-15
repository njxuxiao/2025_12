import pandas as pd
import os
import sys
from openpyxl import load_workbook
import tempfile
import time
import threading
import io
from zipfile import BadZipFile

try:
    if os.name == 'nt':
        import msvcrt
except ImportError:
    pass

try:
    import keyboard
except ImportError:
    print("Error: 'keyboard' library is missing.")
    print("Please run in your terminal or command line: pip install keyboard")
    sys.exit()

try:
    import pykakasi
except ImportError:
    print("Error: 'pykakasi' library is missing.")
    print("Please run in your terminal or command line: pip install pykakasi")
    sys.exit()

try:
    from gtts import gTTS
    import pygame
    gtts_available = True
except ImportError:
    print("Warning: gTTS or pygame library not installed. Online TTS function will be unavailable.")
    print("  - pip install gTTS")
    print("  - pip install pygame")
    gtts_available = False

try:
    import pyttsx3
    pyttsx3_available = True
except ImportError:
    print("Warning: pyttsx3 library not installed. Offline backup TTS function will be unavailable.")
    print("  - pip install pyttsx3")
    pyttsx3_available = False

_gtts_connection_info_printed = False
_gtts_successful_tld = None
_gtts_successful_proxy = None  

tts_control = {
    "engine": None,          
    "stop_requested": False, 
    "lock": threading.Lock() 
}

def get_pyttsx3_japanese_voice_id():
    if not pyttsx3_available:
        return None
    try:
        temp_engine = pyttsx3.init()
        voices = temp_engine.getProperty('voices')
        temp_engine.stop()
        del temp_engine

        japanese_voice_id = None
        for voice in voices:
            lang_str = getattr(voice, 'lang', '').lower()
            name_str = getattr(voice, 'name', '').lower()
            if 'ja' in lang_str or 'japanese' in name_str:
                japanese_voice_id = voice.id
                break
        
        if japanese_voice_id:
            print("\nSuccessfully found [Local] Japanese voice ID. Offline backup is available.")
            return japanese_voice_id
        else:
            print("\nWarning: Japanese [Local] voice package not found in your system. Offline backup may not work.")
            return None
    except Exception as e:
        print(f"\nFailed to scan for local voices: {e}")
        return None


def speak_with_gtts(text):
    global _gtts_connection_info_printed, _gtts_successful_tld, _gtts_successful_proxy
    if not gtts_available or not text:
        return False
    
    success = False
    tld_list = [_gtts_successful_tld] if _gtts_successful_tld else ['com', 'co.jp', 'ca', 'com.au']
    raw_proxy_list = ['http://127.0.0.1:7890', 'http://109.123.97.12:9090']
    
    if _gtts_successful_proxy and _gtts_successful_proxy in raw_proxy_list:
        proxy_list = [_gtts_successful_proxy] + [p for p in raw_proxy_list if p != _gtts_successful_proxy]
    else:
        proxy_list = raw_proxy_list
    
    original_http_proxy = os.environ.get('HTTP_PROXY')
    original_https_proxy = os.environ.get('HTTPS_PROXY')

    try:
        audio_buffer = io.BytesIO()
        
        for proxy_address in proxy_list:
            if success: break 

            os.environ['HTTP_PROXY'] = proxy_address
            os.environ['HTTPS_PROXY'] = proxy_address
            
            if not _gtts_connection_info_printed and proxy_address != _gtts_successful_proxy:
                print(f" -> Attempting connection using Proxy: {proxy_address}")

            for tld in tld_list:
                try:
                    if not _gtts_connection_info_printed:
                        print(f"  -> Trying Google TTS server via tld='{tld}'...")
                    
                    tts = gTTS(text=text, lang='ja', tld=tld)
                    tts.write_to_fp(audio_buffer)
                    success = True
                    
                    _gtts_successful_proxy = proxy_address
                    _gtts_successful_tld = tld
                    
                    if not _gtts_connection_info_printed:
                        print(f"     Success! Connected via '{tld}' using proxy '{proxy_address}'.")
                        _gtts_connection_info_printed = True
                    break 
                except Exception as e:
                    if not _gtts_connection_info_printed:
                        print(f"     Connection to '{tld}' failed: {e}")
                    continue
            
            if not success and not _gtts_connection_info_printed:
                print(f" -> Proxy {proxy_address} failed on all TLDs. Switching to next proxy if available...")

        if success:
            audio_buffer.seek(0)
            pygame.mixer.stop()
            sound = pygame.mixer.Sound(audio_buffer)
            sound.play()
        return success
    except Exception as e:
        print(f"\n!! An unknown error occurred during speech synthesis: {e}")
        return False
    finally:
        if original_http_proxy is None:
            if 'HTTP_PROXY' in os.environ: del os.environ['HTTP_PROXY']
        else:
            os.environ['HTTP_PROXY'] = original_http_proxy
        if original_https_proxy is None:
            if 'HTTPS_PROXY' in os.environ: del os.environ['HTTPS_PROXY']
        else:
            os.environ['HTTPS_PROXY'] = original_https_proxy


def _threaded_speak_pyttsx3(voice_id, text):
    if tts_control["stop_requested"]: return

    try:
        engine = pyttsx3.init()
        
        if tts_control["stop_requested"]:
            engine.stop()
            return

        with tts_control["lock"]:
            if tts_control["stop_requested"]:
                engine.stop()
                return
            tts_control["engine"] = engine
        
        engine.setProperty('voice', voice_id)
        
        if tts_control["stop_requested"]:
            engine.stop()
            return

        engine.say(text)
        engine.runAndWait()
        engine.stop()
        
    except Exception:
        pass 
    finally:
        with tts_control["lock"]:
            tts_control["engine"] = None

def speak_with_pyttsx3(voice_id, text):
    if voice_id and text:
        tts_control["stop_requested"] = False
        if gtts_available:
            pygame.mixer.stop()
            
        t = threading.Thread(target=_threaded_speak_pyttsx3, args=(voice_id, text))
        t.daemon = True
        t.start()


def get_display_length(s):
    length = 0
    for char in str(s):
        if '\u4e00' <= char <= '\u9fff':
            length += 2
        else:
            length += 1
    return length

def display_term(word, grammar):
    print("\n" + "╔" + "═"*50 + "╗")
    print("║" + " "*50 + "║")
    
    if word:
        line = f"  [Word]: {word}"
        padding = 50 - get_display_length(line)
        if padding < 0: padding = 0
        print("║" + line + " "*padding + "║")
        
    if grammar:
        line = f"  [Grammar]: {grammar}"
        padding = 50 - get_display_length(line)
        if padding < 0: padding = 0
        print("║" + line + " "*padding + "║")

    print("║" + " "*50 + "║")
    print("╚" + "═"*50 + "╝")

def display_details(meaning, remarks, word):
    if not meaning and not remarks:
        return

    print("╭" + "┈"*50 + "╮")
    
    if meaning:
        line = f"  [Meaning]: {word}---> {meaning}"
        padding = 50 - get_display_length(line)
        if padding < 0: padding = 0
        print("┆" + line + " "*padding + "┆")

    if remarks:
        line = f"  [Remarks]: {remarks}"
        padding = 50 - get_display_length(line)
        if padding < 0: padding = 0
        print("┆" + line + " "*padding + "┆")

    print("╰" + "┈"*50 + "╯")

def handle_anti_peeking(word, grammar, extra_ui_callback=None):
    tts_control["stop_requested"] = True

    if gtts_available:
        try: pygame.mixer.stop()
        except: pass
    
    with tts_control["lock"]:
        if tts_control["engine"]:
            try: tts_control["engine"].stop()
            except: pass

    os.system('cls' if os.name == 'nt' else 'clear')
    
    print("[]")
    
    time.sleep(0.3)
    
    break_action = False
    
    if os.name == 'nt':
        while msvcrt.kbhit():
            msvcrt.getwch()
        while True:
            key_p = msvcrt.getwch()
            if key_p == '\r':
                break_action = False
                break
            elif key_p.lower() == 'q':
                break_action = True
                break
    else:
        while keyboard.is_pressed('enter'):
            time.sleep(0.05)
        while True:
            event = keyboard.read_event(suppress=False)
            if event.event_type == keyboard.KEY_DOWN:
                if event.name.lower() == 'enter':
                    break_action = False
                    break
                elif event.name.lower() == 'q':
                    break_action = True
                    break

    tts_control["stop_requested"] = False
    
    if not break_action:
        os.system('cls' if os.name == 'nt' else 'clear')
        display_term(word, grammar)
        if extra_ui_callback:
            extra_ui_callback()
            
    return break_action

def smart_sleep(seconds, word, grammar, meaning=None, remarks=None, display_remarks=None):
    start_time = time.time()
    while time.time() - start_time < seconds:
        if keyboard.is_pressed('enter'):
            should_quit = handle_anti_peeking(word, grammar)
            if should_quit:
                return "quit"
            return True 
        time.sleep(0.05)
    return False

def calculate_smart_score(fre, history):
    h = str(history).replace('nan', '')
    fails = h.count('0')
    succs = h.count('1')
    score = fre
    if fails >= 3:
        score += 3
    if succs >= 3:
        score -= 2
    return max(0, score)

def study_helper(file_path, sheet_to_study=None, tts_mode='auto'):
    if gtts_available:
        pygame.init()
        pygame.mixer.init()
    
    try:
        japanese_voice_id = get_pyttsx3_japanese_voice_id()
        auto_mode_current_engine = 'gTTS' if gtts_available else 'pyttsx3'

        kks = pykakasi.kakasi()
        
        all_sheets_data = {}
        chosen_sheet = None

        try:
            if not os.path.exists(file_path):
                print(f"Error: File not found '{file_path}'.")
                return

            print(f"Reading file: {file_path}")
            if os.path.getsize(file_path) == 0:
                print("Error: The file is empty.")
                return

            if file_path.endswith('.xlsx'):
                xls = pd.ExcelFile(file_path, engine='openpyxl')
                sheet_names = xls.sheet_names
                
                if not sheet_names:
                    print("Error: No worksheets found in the Excel file.")
                    return

                if sheet_to_study is not None:
                    if isinstance(sheet_to_study, int) and 0 <= sheet_to_study < len(sheet_names):
                        chosen_sheet = sheet_names[sheet_to_study]
                    elif isinstance(sheet_to_study, str) and sheet_to_study in sheet_names:
                        chosen_sheet = sheet_to_study
                    else:
                        print(f"Error: Specified sheet invalid. Defaulting to first sheet.")
                        chosen_sheet = sheet_names[0]
                else:
                    chosen_sheet = sheet_names[0]
                
                all_sheets_data = pd.read_excel(file_path, sheet_name=None, engine='openpyxl')
                df = all_sheets_data[chosen_sheet]
            else:
                df = pd.read_csv(file_path, sep='\t', encoding='utf-8')
            
            df.columns = df.columns.str.strip()

        except BadZipFile:
            print(f"\nError: Corrupted file. Make sure it is a valid .xlsx file.")
            return
        except Exception as e:
            print(f"Error reading file: {e}")
            return

        if 'Fre' not in df.columns:
            print("No 'Fre' column detected, creating it automatically.")
            df['Fre'] = 0
        else:
            df['Fre'] = pd.to_numeric(df['Fre'], errors='coerce').fillna(0).astype(int)

        if 'History' not in df.columns:
            df['History'] = ""
        else:
            df['History'] = df['History'].fillna("").astype(str).replace('nan', '')

        if '单词' not in df.columns and '文法' not in df.columns:
            print(f"Error: The file must contain at least a '单词' or '文法' column.")
            return
            
        has_reading_col = '读音' in df.columns
        has_meaning_col = '含义' in df.columns
        has_remarks_col = '备注' in df.columns
        
        is_smart_mode = True 

        records = df.to_dict('records')
        original_indices = df.index.tolist()

        zipped = list(zip(records, original_indices))
        zipped.sort(key=lambda x: calculate_smart_score(x[0].get('Fre', 0), x[0].get('History', '')), reverse=True)
        records = [x[0] for x in zipped]
        original_indices = [x[1] for x in zipped]

        print(f"\n--- Japanese Study Helper Started ({'Excel' if file_path.endswith('.xlsx') else 'TXT'} Mode) ---")
        if tts_mode == 'online': print("[TTS Mode]: Online Only")
        elif tts_mode == 'offline': print("[TTS Mode]: Offline Only")
        else: print("[TTS Mode]: Auto")
            
        print("[Mode]: Smart Fre Mode (Press 's' to toggle)")
        print("[IMPORTANT] Please make sure your input method is in English mode.")
        
        is_changed = False
        exited_early = False
        last_answered_correctly_index = None
        
        speak_immediate = False 

        session_history_updated_ids = set()
        session_missed_ids = set()
        session_missed_records = []
        session_missed_original_indices = []
        has_done_bulk_review = False

        def get_fre0_boundary(start_idx, recs):
            for idx in range(start_idx, len(recs)):
                if recs[idx].get('Fre', 0) == 0:
                    return idx
            return len(recs)

        i = 0
        while i < len(records) or (not has_done_bulk_review and session_missed_records):
            
            if i == len(records) and not has_done_bulk_review:
                has_done_bulk_review = True
                if session_missed_records:
                    print("\n[Bulk Review] Reviewing all words missed in this session!")
                    for idx_offset, rec in enumerate(session_missed_records):
                        records.append(rec)
                        original_indices.append(session_missed_original_indices[idx_offset])
                    session_missed_records.clear()
                    
            if i >= len(records):
                break

            current_record = records[i]
            original_index = original_indices[i]

            if not has_done_bulk_review and current_record.get('Fre', 0) == 0:
                has_done_bulk_review = True
                if session_missed_records:
                    print("\n[Bulk Review] Reviewing all words missed in this session before new words!")
                    for idx_offset, rec in enumerate(session_missed_records):
                        records.insert(i + idx_offset, rec)
                        original_indices.insert(i + idx_offset, session_missed_original_indices[idx_offset])
                    session_missed_records.clear()
                    
                    current_record = records[i]
                    original_index = original_indices[i]

            word = str(current_record.get('单词', '')) if pd.notna(current_record.get('单词')) else ""
            grammar = str(current_record.get('文法', '')) if pd.notna(current_record.get('文法')) else ""
            
            reading = ""
            if has_reading_col and pd.notna(current_record.get('读音')):
                reading = str(current_record.get('读音'))
            elif word:
                try:
                    result = kks.convert(word)
                    reading = "".join([item['hira'] for item in result])
                except Exception:
                    pass

            meaning = str(current_record.get('含义', '')) if has_meaning_col and pd.notna(current_record.get('含义')) else ""
            remarks = str(current_record.get('备注', '')) if has_remarks_col and pd.notna(current_record.get('备注')) else ""

            if not word and not grammar:
                i += 1
                continue
            
            display_term(word, grammar)

            text_to_speak = word if word else grammar
            
            def speak():
                if tts_control["stop_requested"]: return

                if tts_mode == 'online':
                    speak_with_gtts(text_to_speak)
                    return

                if tts_mode == 'offline':
                    if japanese_voice_id:
                        speak_with_pyttsx3(japanese_voice_id, text_to_speak)
                    return
                
                nonlocal auto_mode_current_engine
                if auto_mode_current_engine == 'gTTS':
                    speak_with_gtts(text_to_speak)
                elif auto_mode_current_engine == 'pyttsx3':
                    speak_with_pyttsx3(japanese_voice_id, text_to_speak)

            if speak_immediate:
                speak()

            key = None
            while True:
                prompt = "Press a key... (→: Know / 0: Don't Know / q: Quit"
                if last_answered_correctly_index is not None:
                    prompt += " / x: Correct Last / g: Master Last"
                prompt += " / s: Toggle Smart Mode / L: Toggle TTS / Enter: Privacy Mode)" 
                print(prompt + " " + str(i+1) + "/" + str(len(records)), flush=True)
                
                if os.name == 'nt':
                    try:
                        key_stroke = msvcrt.getwch()
                        
                        if key_stroke in ['\xe0', '\x00']:
                            key_stroke = msvcrt.getwch()
                            if key_stroke == 'M': key = 'right'
                            elif key_stroke == 'K': key = 'left'
                            else: key = 'unknown'
                        elif key_stroke == '\r': key = 'enter'
                        elif key_stroke.lower() == 'q': key = 'q'
                        elif key_stroke == '0': key = '0'
                        elif key_stroke.lower() == 'x': key = 'x'
                        elif key_stroke.lower() == 'l': key = 'l'
                        elif key_stroke.lower() == 's': key = 's'
                        elif key_stroke.lower() == 'g': key = 'g'
                        else: key = 'unknown'
                    except Exception:
                        key = 'unknown'
                else:
                    event = keyboard.read_event(suppress=True)
                    while event.event_type != keyboard.KEY_DOWN:
                        event = keyboard.read_event(suppress=True)
                    key = event.name.lower()

                if key == 'l':
                    speak_immediate = not speak_immediate
                    mode_str = "Immediate (Speaking now)" if speak_immediate else "After Answer"
                    print(f"\n[TTS Mode Switched]: {mode_str}")
                    if speak_immediate:
                        speak()
                    continue 

                if key == 's':
                    is_smart_mode = not is_smart_mode
                    mode_str = "Smart Fre Mode" if is_smart_mode else "Normal Mode"
                    print(f"\n[Mode Switched]: {mode_str}")
                    
                    sort_start_idx = i + 1 + 5
                    if sort_start_idx < len(records):
                        to_sort_recs = records[sort_start_idx:]
                        to_sort_orig = original_indices[sort_start_idx:]
                        zipped_rem = list(zip(to_sort_recs, to_sort_orig))
                        
                        if is_smart_mode:
                            zipped_rem.sort(key=lambda x: calculate_smart_score(x[0].get('Fre', 0), x[0].get('History', '')), reverse=True)
                        else:
                            zipped_rem.sort(key=lambda x: x[0].get('Fre', 0), reverse=True)
                        
                        records[sort_start_idx:] = [x[0] for x in zipped_rem]
                        original_indices[sort_start_idx:] = [x[1] for x in zipped_rem]
                        
                    continue

                if key == 'enter':
                    should_quit = handle_anti_peeking(word, grammar)
                    if should_quit:
                        key = 'q' 
                        break
                    continue
                
                if key in ['right', '0', 'q', 'x', 'g']:
                    break
            
            if key == 'q':
                print("Saving progress and exiting early...")
                exited_early = True
                break 
            
            if key == 'x':
                if last_answered_correctly_index is not None:
                    for idx, record in enumerate(records):
                        if original_indices[idx] == last_answered_correctly_index:
                            record['Fre'] += 1
                            
                            h = str(record.get('History', '')).replace('nan', '')
                            if h and h[-1] == '1':
                                record['History'] = h[:-1] + '0'
                            elif not h or h[-1] != '0':
                                record['History'] = (h + '0')[-5:]
                            
                            if id(record) not in session_missed_ids:
                                session_missed_records.append(record)
                                session_missed_original_indices.append(last_answered_correctly_index)
                                session_missed_ids.add(id(record))

                            boundary = get_fre0_boundary(i + 1, records)
                            insert_pos = min(i + 5, boundary)
                            records.insert(insert_pos, record)
                            original_indices.insert(insert_pos, last_answered_correctly_index)
                            print(f"  -> [Spaced Repetition] Word inserted ahead for review (before new words).")
                            
                            break
                    is_changed = True
                    print(f"\nCorrected the previous item!")
                    last_answered_correctly_index = None
                else:
                    print("\nThere is no previous item to correct.")
                continue

            if key == 'g':
                if last_answered_correctly_index is not None:
                    for idx, record in enumerate(records):
                        if original_indices[idx] == last_answered_correctly_index:
                            record['Fre'] = 0
                            break
                    is_changed = True
                    print(f"\nMastered the previous item! Forgotten count reset to 0.")
                    last_answered_correctly_index = None
                else:
                    print("\nThere is no previous item to master.")
                continue

            last_answered_correctly_index = None
            display_remarks = remarks if remarks else reading

            if key == '0':
                current_record['Fre'] += 1
                
                if id(current_record) not in session_history_updated_ids:
                    h = str(current_record.get('History', '')).replace('nan', '')
                    current_record['History'] = (h + "0")[-5:]
                    session_history_updated_ids.add(id(current_record))
                
                is_changed = True
                print(f"Recorded! Forgotten count: {current_record['Fre']}")
                display_details(meaning, display_remarks, word if word else grammar)
                
                if not speak_immediate:
                    speak() 
                
                if id(current_record) not in session_missed_ids:
                    session_missed_records.append(current_record)
                    session_missed_original_indices.append(original_index)
                    session_missed_ids.add(id(current_record))

                boundary = get_fre0_boundary(i + 1, records)
                insert_pos = min(i + 5, boundary)
                records.insert(insert_pos, current_record)
                original_indices.insert(insert_pos, original_index)
                print(f"  -> [Spaced Repetition] Will review again shortly (before new words).")
                
                res = smart_sleep(2.0, word, grammar, meaning, display_remarks)
                if res == "quit":
                    print("Saving progress and exiting early...")
                    exited_early = True
                    break

            elif key == 'right':
                if id(current_record) not in session_history_updated_ids:
                    h = str(current_record.get('History', '')).replace('nan', '')
                    current_record['History'] = (h + "1")[-5:]
                    session_history_updated_ids.add(id(current_record))
                    
                is_changed = True
                
                display_details(meaning, display_remarks, word)
                print(f"Great! Forgotten count: {current_record['Fre']}")
                last_answered_correctly_index = original_index
                
                if not speak_immediate:
                    speak()
                
                res = smart_sleep(0.3, word, grammar, meaning, display_remarks)
                if res == "quit":
                    print("Saving progress and exiting early...")
                    exited_early = True
                    break

            i += 1

        if not exited_early:
            print("\nAll words/grammar have been studied!")
        else:
            print("\nSession ended early.")

        if not is_changed:
            print("\nNo changes were made, no need to save.")
            return
            
        try:
            print("Updating records back to DataFrame...")
            
            unique_records = []
            seen_ids = set()
            for rec in records:
                if id(rec) not in seen_ids:
                    unique_records.append(rec)
                    seen_ids.add(id(rec))
            
            new_df = pd.DataFrame(unique_records) 
            
            if file_path.endswith('.xlsx'):
                all_sheets_data[chosen_sheet] = new_df
                
                print("Saving file and preserving column widths...")
                book = load_workbook(file_path)
                col_widths = {}
                for sheet_name in book.sheetnames:
                    col_widths[sheet_name] = {
                        letter: dim.width for letter, dim in book[sheet_name].column_dimensions.items()
                    }
                
                with pd.ExcelWriter(file_path, engine='openpyxl') as writer:
                    for sheet_name, sheet_data in all_sheets_data.items():
                        sheet_data.to_excel(writer, sheet_name=sheet_name, index=False)
                        
                        if sheet_name in col_widths:
                            ws = writer.sheets[sheet_name]
                            for col_letter, width in col_widths[sheet_name].items():
                                if width:
                                   ws.column_dimensions[col_letter].width = width
            else:
                new_df.to_csv(file_path, sep='\t', encoding='utf-8', index=False)
            
            print("\nStudy session finished! Your progress has been saved successfully.")
        
        except Exception as e:
            print(f"\nAn unknown error occurred while saving the file: {e}")
    finally:
        if gtts_available:
            pygame.quit()


if __name__ == '__main__':
    
    excel_list = ["0_transform.xlsx", "1_zhouyf.xlsx"]

    excel_file_path = excel_list[1]
    
    study_sheet = 2

    preferred_tts_engine = 'auto' 

    study_helper(excel_file_path, sheet_to_study=study_sheet, tts_mode=preferred_tts_engine)

    for i in range(10): print("\n")